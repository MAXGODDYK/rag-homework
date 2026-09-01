from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from jarvis.ingestion.archives import ArchiveSafetyError, extract_archive
from jarvis.ingestion.parsers import parse_document
from jarvis.ingestion.repository import analyze_code
from jarvis.ingestion.service import IngestionService
from jarvis.ingestion.vector_index import VectorIndex
from jarvis.retrieval import DynamicRetriever

from helpers import make_config, make_database


def test_docx_heading_and_text_are_preserved(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_heading("Plan", level=1)
    document.add_paragraph("Prepare the first topic.")
    document.save(path)

    parsed = parse_document(path)
    assert [unit.text for unit in parsed.units] == ["Plan", "Prepare the first topic."]
    assert parsed.units[1].heading == "Plan"


def test_xlsx_citation_uses_real_column_letters(tmp_path: Path) -> None:
    path = tmp_path / "wide.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["AA2"] = "value"
    workbook.save(path)

    parsed = parse_document(path)
    assert parsed.units[0].sheet == "Sheet"
    assert parsed.units[0].cell_range == "A1:AA2"


def test_zip_traversal_is_rejected(tmp_path: Path) -> None:
    archive = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(archive, "w") as stream:
        stream.writestr("../escape.txt", "not allowed")

    with pytest.raises(ArchiveSafetyError):
        extract_archive(archive, tmp_path / "out", max_files=10, max_extracted_bytes=1024)
    assert not (tmp_path / "escape.txt").exists()


def test_nested_archive_is_not_expanded(tmp_path: Path) -> None:
    nested_bytes = io.BytesIO()
    with zipfile.ZipFile(nested_bytes, "w") as nested:
        nested.writestr("inside.txt", "nested")
    archive = tmp_path / "outer.zip"
    with zipfile.ZipFile(archive, "w") as outer:
        outer.writestr("readme.txt", "safe")
        outer.writestr("nested.zip", nested_bytes.getvalue())

    extracted = extract_archive(
        archive, tmp_path / "out", max_files=10, max_extracted_bytes=1024 * 1024
    )
    assert [path.name for path in extracted] == ["readme.txt"]
    assert not (tmp_path / "out" / "nested.zip").exists()


def test_repository_fallback_extracts_symbols_and_imports(tmp_path: Path) -> None:
    path = tmp_path / "module.py"
    text = "import json\n\nclass Service:\n    def run(self):\n        return json.dumps({})\n"
    symbols, edges = analyze_code(path, text)
    names = {symbol.name for symbol in symbols}
    assert {"Service", "run"}.issubset(names)
    assert any(edge.target_ref == "json" for edge in edges)


@pytest.mark.parametrize(
    ("name", "content", "expected"),
    [
        ("notes.txt", "plain text", "plain text"),
        ("notes.md", "# Heading\nbody", "Heading"),
        ("data.json", '{"topic":"planning"}', "planning"),
        ("data.yaml", "topic: planning", "planning"),
        ("table.csv", "topic,value\nplanning,1", "planning"),
        ("page.html", "<html><body><h1>Plan</h1><p>Study</p></body></html>", "Study"),
    ],
)
def test_text_family_parsers(tmp_path: Path, name: str, content: str, expected: str) -> None:
    path = tmp_path / name
    path.write_text(content, encoding="utf-8")
    parsed = parse_document(path)
    assert expected in "\n".join(unit.text for unit in parsed.units)


def test_notebook_cells_are_separate_units(tmp_path: Path) -> None:
    path = tmp_path / "analysis.ipynb"
    path.write_text(
        json.dumps(
            {
                "cells": [
                    {"cell_type": "markdown", "source": ["# Result"]},
                    {"cell_type": "code", "source": ["print(42)"]},
                ]
            }
        ),
        encoding="utf-8",
    )
    parsed = parse_document(path)
    assert [unit.metadata["cell_type"] for unit in parsed.units] == ["markdown", "code"]


def test_pptx_slide_number_is_preserved(tmp_path: Path) -> None:
    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Architecture"
    slide.placeholders[1].text = "Backend"
    presentation.save(path)
    parsed = parse_document(path)
    assert parsed.units[0].page == 1
    assert "Backend" in parsed.units[0].text


def test_pdf_without_text_layer_requests_ocr(tmp_path: Path) -> None:
    from pypdf import PdfWriter

    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with path.open("wb") as stream:
        writer.write(stream)
    with pytest.raises(Exception, match="OCR plugin is required"):
        parse_document(path)


def test_ingestion_is_isolated_by_user_and_project(monkeypatch, tmp_path: Path) -> None:
    config = make_config(tmp_path)
    database = make_database(config)
    database.ensure_user("other", "Other")
    project = database.create_project("owner", "Owner project", None)
    other_project = database.create_project("other", "Other project", None)
    service = IngestionService(config, database)
    monkeypatch.setattr(VectorIndex, "rebuild", lambda self, ids, texts: None)
    monkeypatch.setattr(VectorIndex, "search", lambda self, query, top_k: [])

    owner_file = tmp_path / "owner.txt"
    owner_file.write_text("unique alpha planning phrase", encoding="utf-8")
    other_file = tmp_path / "other.txt"
    other_file.write_text("unique beta private phrase", encoding="utf-8")
    owner_record = service.ingest_upload(
        owner_file, user_id="owner", project_id=project["id"]
    )[0]
    service.ingest_upload(
        other_file, user_id="other", project_id=other_project["id"]
    )

    retriever = DynamicRetriever(config, database)
    owner_rows = retriever.search(
        "alpha planning", user_id="owner", project_id=project["id"], rerank=False
    )
    leaked_rows = retriever.search(
        "beta private", user_id="owner", project_id=project["id"], rerank=False
    )
    assert owner_rows and all(row["user_id"] == "owner" for row in owner_rows)
    assert leaked_rows == []

    service.delete_document(owner_record["document_id"], user_id="owner")
    assert database.query_one(
        "SELECT id FROM documents WHERE id=?", (owner_record["document_id"],)
    ) is None
