from __future__ import annotations

import csv
import io
import json
import mimetypes
import re
import zipfile
from pathlib import Path
from typing import Callable

from bs4 import BeautifulSoup
from charset_normalizer import from_bytes

from .models import ExtractedUnit, ParsedDocument


class DocumentParseError(RuntimeError):
    """Safe unsupported or malformed document failure."""


CODE_EXTENSIONS = frozenset(
    {
        ".py", ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx", ".css", ".scss",
        ".vue", ".svelte", ".cs", ".java", ".kt", ".kts", ".go", ".rs", ".c",
        ".h", ".cc", ".cpp", ".hpp", ".php", ".rb", ".swift", ".dart", ".lua",
        ".sql", ".ps1", ".sh", ".bash", ".zsh", ".proto", ".gradle", ".fs", ".fsx",
    }
)
TEXT_EXTENSIONS = CODE_EXTENSIONS | frozenset(
    {
        ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".jsonl",
        ".xml", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".html", ".htm",
        ".ipynb", ".dockerfile", ".makefile",
    }
)
BLOCKED_FILENAMES = frozenset(
    {".env", "credentials", "credentials.json", "id_rsa", "id_ed25519", "secrets.yml"}
)
BLOCKED_EXTENSIONS = frozenset(
    {".exe", ".dll", ".sys", ".msi", ".com", ".scr", ".bat", ".cmd", ".pfx", ".pem", ".key"}
)


def is_blocked_path(path: Path) -> bool:
    lowered = path.name.lower()
    return lowered in BLOCKED_FILENAMES or path.suffix.lower() in BLOCKED_EXTENSIONS


def is_probably_text(path: Path) -> bool:
    if is_blocked_path(path):
        return False
    if path.suffix.lower() in TEXT_EXTENSIONS or path.name.lower() in {"dockerfile", "makefile"}:
        return True
    try:
        sample = path.read_bytes()[:8192]
    except OSError:
        return False
    return b"\x00" not in sample and bool(from_bytes(sample).best())


def _decode(path: Path) -> str:
    data = path.read_bytes()
    if b"\x00" in data[:8192]:
        raise DocumentParseError("Binary content cannot be parsed as text")
    best = from_bytes(data).best()
    if best is None:
        raise DocumentParseError("Text encoding could not be detected")
    return str(best)


def _plain(path: Path) -> ParsedDocument:
    text = _decode(path)
    return ParsedDocument(
        path=path,
        media_type=mimetypes.guess_type(path.name)[0] or "text/plain",
        units=[ExtractedUnit(text=text, line_start=1, line_end=max(1, len(text.splitlines())))],
        metadata={"kind": "code" if path.suffix.lower() in CODE_EXTENSIONS else "text"},
    )


def _html(path: Path) -> ParsedDocument:
    soup = BeautifulSoup(_decode(path), "html.parser")
    for element in soup(["script", "style", "noscript", "template"]):
        element.decompose()
    units: list[ExtractedUnit] = []
    current_heading: str | None = None
    for element in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "code", "table"]):
        text = " ".join(element.stripped_strings)
        if not text:
            continue
        if element.name and element.name.startswith("h"):
            current_heading = text
        units.append(ExtractedUnit(text=text, heading=current_heading, metadata={"tag": element.name}))
    return ParsedDocument(path=path, media_type="text/html", units=units, metadata={"title": soup.title.string if soup.title and soup.title.string else None})


def _pdf(path: Path) -> ParsedDocument:
    from pypdf import PdfReader

    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise DocumentParseError("Encrypted PDFs are not supported")
        units = [
            ExtractedUnit(text=page.extract_text() or "", page=index)
            for index, page in enumerate(reader.pages, start=1)
        ]
    except DocumentParseError:
        raise
    except Exception as error:
        raise DocumentParseError("PDF could not be parsed") from error
    if not any(unit.text.strip() for unit in units):
        raise DocumentParseError("PDF has no text layer; OCR plugin is required")
    return ParsedDocument(path=path, media_type="application/pdf", units=units, metadata={"pages": len(units)})


def _docx(path: Path) -> ParsedDocument:
    from docx import Document

    try:
        document = Document(path)
        units: list[ExtractedUnit] = []
        heading: str | None = None
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                heading = text
            units.append(ExtractedUnit(text=text, heading=heading))
        for table_index, table in enumerate(document.tables, start=1):
            rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            units.append(ExtractedUnit(text="\n".join(rows), heading=f"Table {table_index}"))
    except Exception as error:
        raise DocumentParseError("DOCX could not be parsed") from error
    return ParsedDocument(path=path, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", units=units)


def _pptx(path: Path) -> ParsedDocument:
    from pptx import Presentation

    try:
        presentation = Presentation(path)
        units = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            text = "\n".join(
                shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()
            )
            units.append(ExtractedUnit(text=text, page=slide_number, heading=f"Slide {slide_number}"))
    except Exception as error:
        raise DocumentParseError("PPTX could not be parsed") from error
    return ParsedDocument(path=path, media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation", units=units)


def _xlsx(path: Path) -> ParsedDocument:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        units = []
        for sheet in workbook.worksheets:
            rows: list[str] = []
            max_row = 0
            max_column = 0
            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                values = ["" if cell.value is None else str(cell.value) for cell in row]
                if any(values):
                    rows.append("\t".join(values))
                    max_row = max(max_row, row_number)
                    last_value_column = max(
                        index for index, value in enumerate(values, start=1) if value
                    )
                    max_column = max(max_column, last_value_column)
            if rows:
                end_column = get_column_letter(max_column)
                units.append(
                    ExtractedUnit(
                        text="\n".join(rows),
                        sheet=sheet.title,
                        cell_range=f"A1:{end_column}{max_row}",
                    )
                )
        workbook.close()
    except Exception as error:
        raise DocumentParseError("XLSX could not be parsed") from error
    return ParsedDocument(path=path, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", units=units)


def _odf(path: Path) -> ParsedDocument:
    from odf import teletype
    from odf.opendocument import load
    from odf.text import H, P

    try:
        document = load(str(path))
        units = []
        heading: str | None = None
        for node in document.getElementsByType(H) + document.getElementsByType(P):
            text = teletype.extractText(node).strip()
            if not text:
                continue
            if node.qname[1] == "h":
                heading = text
            units.append(ExtractedUnit(text=text, heading=heading))
    except Exception as error:
        raise DocumentParseError("OpenDocument file could not be parsed") from error
    return ParsedDocument(path=path, media_type=mimetypes.guess_type(path.name)[0] or "application/vnd.oasis.opendocument", units=units)


def _epub(path: Path) -> ParsedDocument:
    from ebooklib import ITEM_DOCUMENT, epub

    try:
        book = epub.read_epub(str(path))
        units = []
        for item in book.get_items_of_type(ITEM_DOCUMENT):
            soup = BeautifulSoup(item.get_content(), "html.parser")
            for element in soup(["script", "style"]):
                element.decompose()
            text = "\n".join(soup.stripped_strings)
            if text:
                units.append(ExtractedUnit(text=text, heading=item.get_name()))
    except Exception as error:
        raise DocumentParseError("EPUB could not be parsed") from error
    return ParsedDocument(path=path, media_type="application/epub+zip", units=units)


def _notebook(path: Path) -> ParsedDocument:
    try:
        payload = json.loads(_decode(path))
        units = []
        for index, cell in enumerate(payload.get("cells", []), start=1):
            source = "".join(cell.get("source", []))
            if source.strip():
                units.append(
                    ExtractedUnit(
                        text=source,
                        heading=f"{cell.get('cell_type', 'unknown')} cell {index}",
                        metadata={"cell_type": cell.get("cell_type")},
                    )
                )
    except Exception as error:
        raise DocumentParseError("Jupyter notebook could not be parsed") from error
    return ParsedDocument(path=path, media_type="application/x-ipynb+json", units=units)


PARSERS: dict[str, Callable[[Path], ParsedDocument]] = {
    ".pdf": _pdf,
    ".docx": _docx,
    ".pptx": _pptx,
    ".xlsx": _xlsx,
    ".odt": _odf,
    ".ods": _odf,
    ".odp": _odf,
    ".epub": _epub,
    ".html": _html,
    ".htm": _html,
    ".ipynb": _notebook,
}


def parse_document(path: Path) -> ParsedDocument:
    if is_blocked_path(path):
        raise DocumentParseError(f"Blocked file type: {path.name}")
    parser = PARSERS.get(path.suffix.lower())
    if parser:
        return parser(path)
    if is_probably_text(path):
        return _plain(path)
    raise DocumentParseError(f"Unsupported or binary file: {path.name}")
